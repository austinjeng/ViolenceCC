# -*- coding: utf-8 -*-
"""Build the CGW '26 talk deck (paper #32) from paper assets.

Run in the vcc-main env from the repo root:
    conda run -n vcc-main python scripts/build_talk_pptx.py

Produces:
    presentation/assets/*.png   (figure conversions + rendered equations)
    presentation/32_Dual-Modal Skeleton-Visual Fusion for Weakly Supervised
        Violence Detection：A Multi-Backbone Study.pptx   (fullwidth colon)

Content source of truth: presentation/slides.md. All numbers are 3-seed
means from paper/main.tex Tables 1-3.
"""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import fitz  # pymupdf
from PIL import Image

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "presentation" / "assets"
FIGDIR = ROOT / "paper" / "figures"
OUT = ROOT / "presentation" / (
    "32_Dual-Modal Skeleton-Visual Fusion for Weakly Supervised "
    "Violence Detection：A Multi-Backbone Study.pptx"
)

# ---- palette (paper Fig. 1) -------------------------------------------------
INK = RGBColor(0x2C, 0x3E, 0x50)
MUTE = RGBColor(0x7F, 0x8C, 0x8D)
SOFT = RGBColor(0xF4, 0xF6, 0xF7)
GREEN = RGBColor(0x1E, 0x84, 0x49)
GREEN_L = RGBColor(0xD5, 0xF5, 0xE3)
BLUE = RGBColor(0x24, 0x71, 0xA3)
BLUE_L = RGBColor(0xEB, 0xF5, 0xFB)
ORANGE = RGBColor(0xD6, 0x89, 0x10)
ORANGE_L = RGBColor(0xFE, 0xF9, 0xE7)
RED = RGBColor(0xC0, 0x39, 0x2B)
PURPLE = RGBColor(0x7D, 0x3C, 0x98)
PURPLE_L = RGBColor(0xF5, 0xEE, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"

SW, SH = 13.333, 7.5  # slide size in inches

# ---- asset generation -------------------------------------------------------

FIG_ZOOM = {
    "fig_backbone_comparison": 6,
    "fig_gating_distribution": 6,
    "fig_temporal_scores": 8,
    "fig_tta_comparison": 8,
}


def convert_figures():
    ASSETS.mkdir(parents=True, exist_ok=True)
    for name, zoom in FIG_ZOOM.items():
        src = FIGDIR / f"{name}.pdf"
        dst = ASSETS / f"{name}.png"
        doc = fitz.open(src)
        pix = doc[0].get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        pix.save(dst)
        doc.close()
        print(f"  converted {name}.png ({pix.width}x{pix.height})")


EQUATIONS = {
    "eq_proj": r"$\hat{\mathbf{s}} = \mathrm{LN}_s(W_s\,\mathbf{s} + \mathbf{b}_s)"
               r" \qquad \hat{\mathbf{v}} = \mathrm{LN}_v(W_v\,\mathbf{v} + \mathbf{b}_v)$",
    "eq_gate": r"$\mathbf{g} = \sigma\left(W_g\,[\,\hat{\mathbf{s}}\,\Vert\,\hat{\mathbf{v}}\,]"
               r" + \mathbf{b}_g\right)$",
    "eq_fuse": r"$\mathbf{f} = \mathrm{LN}_f\left(\mathbf{g}\odot\hat{\mathbf{s}}"
               r" + (\mathbf{1}-\mathbf{g})\odot\hat{\mathbf{v}}"
               r" + \hat{\mathbf{s}} + \hat{\mathbf{v}}\right)$",
    "eq_mil": r"$\mathcal{L} = \max\left(0,\ 1 - \frac{1}{k}\sum_{i\in\mathcal{A}_k} a_i"
              r" + \frac{1}{k}\sum_{j\in\mathcal{N}_k} a_j\right)"
              r" + \lambda_1\,\Omega_{\mathrm{sparse}}"
              r" + \lambda_2 \sum_i (a_i - a_{i+1})^2$",
    "eq_w": r"$w = 1-(1-w_{\mathrm{vl}})\,s \qquad"
            r" w_{\mathrm{vl}} = \min\left(1,\ \sigma_{\mathrm{test}}/\sigma_{\mathrm{clean}}\right)$",
}


def render_equations():
    plt.rcParams["mathtext.fontset"] = "cm"
    for name, tex in EQUATIONS.items():
        fig = plt.figure(figsize=(10, 1.2))
        fig.patch.set_alpha(0)
        fig.text(0.5, 0.5, tex, fontsize=22, ha="center", va="center",
                 color="#2C3E50")
        fig.savefig(ASSETS / f"{name}.png", dpi=300, transparent=True,
                    bbox_inches="tight", pad_inches=0.06)
        plt.close(fig)
        print(f"  rendered {name}.png")


def render_temporal_talk_figure():
    """Slide-specific qualitative example.

    The paper's Figure 2 video (RoadAccidents127) is anti-aligned with its
    ground truth (worst fusion advantage among all 128 annotated anomalous
    test videos, 128/128), so the talk uses Fighting047 instead: gated-fusion
    GT-vs-normal score gap +0.48 vs +0.15 (CLIP-only) and -0.01
    (skeleton-only), seed-42 runs.
    allow_pickle=True is safe: eval_scores.npz files are produced by this
    repo's own pipeline (same pattern as scripts/generate_pub_figures.py).
    """
    vid = "Fighting047"
    ann = {}
    for line in (ROOT / "data/annotations/ucf_temporal.txt").read_text() \
            .splitlines():
        parts = line.split()
        if len(parts) >= 6:
            name = parts[0].replace("_x264.mp4", "")
            iv = [(int(parts[2]), int(parts[3]))]
            if int(parts[4]) != -1:
                iv.append((int(parts[4]), int(parts[5])))
            ann[name] = iv
    runs = {
        "Skeleton Only": ("ucf_skeleton_only_s42",
                          dict(color="#7F8C8D", ls="--", lw=2.2)),
        "Visual Only (CLIP)": ("ucf_clip_only_s42",
                               dict(color="#2C3E50", ls="-.", lw=2.2)),
        "Gated Fusion": ("ucf_gated_fusion_s42",
                         dict(color="#C0392B", ls="-", lw=2.8)),
    }
    plt.rcParams.update({"font.size": 15, "font.family": "serif"})
    fig, ax = plt.subplots(figsize=(9.2, 3.9))
    first = True
    for a, b in ann[vid]:
        ax.axvspan(a, b, color="#C0392B", alpha=0.14,
                   label="Ground Truth" if first else None)
        first = False
    for label, (run, style) in runs.items():
        d = np.load(ROOT / f"results/{run}/eval_scores.npz",
                    allow_pickle=True)
        s = np.asarray(d[vid], dtype=float)
        ax.plot(np.arange(len(s)), s, label=label, **style)
    ax.set_xlabel("Frame Number")
    ax.set_ylabel("Anomaly Score")
    ax.set_title(f"{vid} (Fighting)")
    ax.set_ylim(-0.02, 1.05)
    ax.set_xlim(0, None)
    ax.grid(alpha=0.25)
    ax.legend(loc="center right", framealpha=0.95)
    fig.tight_layout()
    fig.savefig(ASSETS / "fig_temporal_scores_talk.png", dpi=220)
    plt.close(fig)
    print("  rendered fig_temporal_scores_talk.png (Fighting047)")


def png_size_in(path, scale=1.30, dpi=300):
    """Displayed (w, h) in inches for an equation PNG at a consistent scale."""
    with Image.open(path) as im:
        return im.width / dpi * scale, im.height / dpi * scale


# ---- deck helpers -----------------------------------------------------------

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_strip(slide):
    """4-color modality strip along the top edge."""
    colors = [GREEN, BLUE, ORANGE, PURPLE]
    seg = SW / 4
    for i, c in enumerate(colors):
        r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(i * seg),
                                   Inches(0), Inches(seg), Inches(0.06))
        r.fill.solid()
        r.fill.fore_color.rgb = c
        r.line.fill.background()
        r.shadow.inherit = False


def add_text(slide, x, y, w, h, runs_list, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=6, line_spacing=1.0):
    """runs_list: list of paragraphs; each paragraph is a list of
    (text, size_pt, bold, color) run tuples, or a dict with extra keys."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs_list):
        opts = {}
        if isinstance(para, dict):
            opts = para
            para = para["runs"]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        p.space_before = Pt(opts.get("space_before", 0))
        if line_spacing != 1.0:
            p.line_spacing = line_spacing
        for text, size, bold, color in para:
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
    return box


def add_title(slide, text, size=27):
    add_text(slide, 0.45, 0.22, SW - 0.9, 0.75,
             [[(text, size, True, INK)]])


def add_footer(slide, num):
    add_text(slide, 0.45, SH - 0.36, 8.5, 0.3,
             [[("CGW '26 · Paper #32 · Dual-Modal Skeleton-Visual "
                "Fusion for Weakly Supervised Violence Detection", 8, False,
                MUTE)]])
    add_text(slide, SW - 1.0, SH - 0.36, 0.6, 0.3,
             [[(str(num), 9, False, MUTE)]], align=PP_ALIGN.RIGHT)


def add_box(slide, x, y, w, h, fill, line_color, text_paras=None, radius=0.10,
            line_w=1.25, dash=None, anchor=MSO_ANCHOR.MIDDLE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                 Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = dash
    if text_paras:
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.07)
        tf.margin_top = tf.margin_bottom = Inches(0.04)
        for i, para in enumerate(text_paras):
            opts = {}
            if isinstance(para, dict):
                opts = para
                para = para["runs"]
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = opts.get("align", PP_ALIGN.CENTER)
            p.space_after = Pt(opts.get("space_after", 2))
            for text, size, bold, color in para:
                r = p.add_run()
                r.text = text
                r.font.name = FONT
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.color.rgb = color
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=INK, width=1.75, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                      Inches(y1), Inches(x2), Inches(y2))
    conn.shadow.inherit = False
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash:
        conn.line.dash_style = dash
    ln = conn.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return conn


def set_cell(cell, text, size=10, bold=False, color=INK, fill=WHITE,
             align=PP_ALIGN.CENTER):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.04)
    cell.margin_top = cell.margin_bottom = Inches(0.015)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def make_table(slide, x, y, w, h, n_rows, n_cols, col_widths=None):
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    table = gf.table
    table.first_row = False
    table.horz_banding = False
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    return table


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_picture_fit(slide, path, x, y, w=None, h=None, center_x=None):
    if center_x is not None and w is not None:
        x = center_x - w / 2
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def add_eq(slide, name, x, y, scale=1.30, center_x=None, max_w=None):
    path = ASSETS / f"{name}.png"
    w, h = png_size_in(path, scale=scale)
    if max_w is not None and w > max_w:
        h *= max_w / w
        w = max_w
    if center_x is not None:
        x = center_x - w / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    return h


# ---- speaker notes (Traditional Chinese transcript) -------------------------

NOTES = {
    1: "〔0:00–0:20〕各位老師、各位先進，大家好。我是臺灣科技大學的〔請自行帶入中文姓名〕，"
       "指導教授是楊傳凱教授。今天報告的題目是：結合骨架與視覺語言特徵的雙模態融合，"
       "應用於弱監督暴力偵測，以及跨骨幹網路的系統性研究。",
    2: "〔0:20–1:05〕先從問題背景說起。監視器影像的暴力偵測，最大的瓶頸在標註成本："
       "整段影片標一個「有沒有暴力」的標籤很便宜，但標到每個影格幾乎不可行，"
       "所以主流是弱監督式作法，用 Multiple Instance Learning 只靠影片層級標籤學習。"
       "但現有方法多半只用單一視覺模態——例如 CLIP 特徵——只看外觀；"
       "而打架、攻擊這類暴力事件，本質上是「動作」。"
       "同時，實際部署還會遇到雜訊、模糊、壓縮等畫質劣化問題。",
    3: "〔1:05–1:50〕我們看到三個缺口：第一，骨架動態與視覺語言特徵的融合，"
       "在弱監督 VAD 裡幾乎沒有被探索；第二，融合對視覺骨幹的選擇有多敏感，"
       "沒有系統性研究；第三，TENT、SAR 這類 entropy minimization 的 test-time "
       "adaptation 是為 BatchNorm 設計的，在 LayerNorm 架構上行不行，沒人驗證過。"
       "對應的就是本文三個貢獻：雙模態閘控融合、四種骨幹的系統性比較，"
       "以及 TTA 的深入分析與我們提出的 reliability reweighting。",
    4: "〔1:50–3:05〕這是整體架構。輸入影片走兩條完全凍結的特徵路徑。"
       "上面是骨架流：先用 RTMPose 抽出每幀信心值最高的兩個人、十七個關鍵點，"
       "送進在 NTU RGB+D 120 預訓練的 CTR-GCN，對每個 64 幀片段得到 256 維骨架特徵。"
       "下面是視覺流：約每秒取一張影格，送進凍結的 CLIP 或 SigLIP2 骨幹，"
       "再用 mean 加 max pooling 聚合成片段特徵。"
       "兩條特徵進到中間橘色的閘控融合模組——這是全模型唯一需要訓練的部分，"
       "加上 MIL head 只有五十萬到一百萬個參數，在一張 RTX 4090 上訓練一組設定不到五分鐘。"
       "最後 MIL head 對每個片段輸出異常分數，以 MIL ranking loss 訓練。",
    5: "〔3:05–3:55〕閘控融合的細節：兩個模態各自線性投影到共享的 256 維空間，"
       "各接一個 LayerNorm；接著串接後用 sigmoid 產生逐維度的 gate，"
       "對兩模態做加權混合，再加上殘差保留雙方資訊，最後再過一層 LayerNorm。"
       "訓練用標準 MIL ranking loss：從異常袋與正常袋各取 top-k 分數拉開間距，"
       "再加上稀疏與時間平滑正則項。我們也保留固定等權重的 late fusion 作為對照。",
    6: "〔3:55–4:30〕實驗設計：兩個基準資料集——UCF-Crime 一千九百部影片、"
       "以 frame-level AUC 評估；XD-Violence 四千七百多部、以 AP 評估。"
       "視覺骨幹比較四種：CLIP ViT-B/16 和三種 SigLIP2，從 Base、SO400M 到 Giant。"
       "所有實驗固定三個隨機種子，回報平均加減標準差，切分與流程完全一致。",
    7: "〔4:30–5:45〕主要結果。先看最上面：骨架單獨其實不強，UCF 只有 68.8 的 AUC、"
       "XD 只有 40.8 的 AP，遠低於視覺單模態。但重點是融合：閘控融合在全部八個"
       "「骨幹×資料集」組合裡，全都大於等於視覺單模態，其中六個嚴格更好——"
       "增益不大，平均約 0.6 個百分點，但方向非常一致。"
       "最好的結果：UCF-Crime 用 SigLIP2 Giant 到 82.5% AUC；"
       "XD-Violence 用 SigLIP2 SO400M 到 78.7% AP，這是我們的 headline。"
       "也請注意 late fusion：固定等權重常常反而比視覺單模態差，"
       "在 XD 上與閘控融合差距超過十個 AP 百分點。",
    8: "〔5:45–6:40〕把結果畫成圖更清楚。三個觀察：第一，骨架的互補性小而穩定，"
       "八組全部不退步；第二，閘控真的重要——它會依樣本動態決定信任哪個模態，"
       "late fusion 做不到；第三，骨幹偏好跟資料集有交互作用：UCF 偏好容量最大的 Giant，"
       "而 XD 上四個骨幹的視覺單模態表現在種子變異內難分高下，"
       "是閘控融合之後 SO400M 才明顯勝出。這提醒我們：挑骨幹不能只看通用 benchmark。",
    9: "〔6:40–7:05〕一個定性例子：UCF-Crime 的打架影片。紅色實線是閘控融合的分數："
       "整段打鬥期間都維持高分——包括視覺單模態一開始漏掉的前段——事件結束後降回低分；"
       "骨架單獨則幾乎沒有反應。",
    10: "〔7:05–7:55〕跟文獻比較，我們刻意誠實：在「凍結特徵、不用文字對齊、單卡可訓」"
        "的同級方法裡，XD 的 78.7 AP 與 RTFM、MGFN 相當；UCF 的 82.5 則落後這個級距"
        "所有的現代凍結特徵方法，我們不主張 SOTA。另一個值得注意的點：training-free 不等於便宜——"
        "EventVAD 用 80GB 的 A800、LAVAD 用雙張 RTX 3090，"
        "我們在單張 4090 上 UCF 持平或更好、XD 領先超過 14 個 AP。",
    11: "〔7:55–9:05〕最後是 test-time adaptation。我們在 UCF-Crime-C 上測試："
        "四種劣化、五個嚴重度、共二十個條件、三個種子。第一個發現：TENT 和 SAR "
        "只能更新融合頭裡 LayerNorm 的一千五百多個 affine 參數，"
        "結果每個骨幹的 AUC 變化都小於 0.1 個百分點——等於零。這是結構性的："
        "分布偏移發生在凍結的骨幹特徵，而 AUC 是純排序指標。"
        "所以我們換個思路，提出 discriminative-reliability reweighting："
        "完全不更新參數，監測視覺流分數的離散度，一旦崩塌，"
        "就把融合導向較耐劣化的骨架流。平均提升 1.21 個百分點，四個骨幹、"
        "每個種子全部為正；在視覺流崩最嚴重的高斯雜訊最高嚴重度可達 +13.2；"
        "其他劣化型態增益為零——代表 gate 正確地選擇不出手。"
        "這是「救援」而不是全面的 robustness。",
    12: "〔9:05–9:45〕總結三點：第一，骨架特徵雖弱，卻能對強的視覺骨幹提供小而一致的"
        "互補增益，而且不會造成退步；第二，骨幹選擇與資料集有交互作用，且透過融合才顯現，"
        "多骨幹評估有其必要；第三，entropy TTA 只能動 LN 的 affine 參數，"
        "在我們這種凍結的排序式偵測器上推不動；reliability reweighting 可以，"
        "前提是至少一個模態仍然可靠。"
        "限制包括凍結設計的效能上限、reweighting 的 transductive 統計假設，"
        "以及只涵蓋兩個資料集與合成劣化。",
    13: "〔9:45–10:00〕以上是今天的報告，程式碼與實驗設定將會公開。謝謝大家，歡迎提問。",
    14: "備用投影片，供 Q&A 使用。",
    15: "〔Q&A 備用〕若被問到與 SOTA 的完整比較：上半是重預算方法（文字對齊、額外模態、"
        "大型語言模型），下半是與我們同級的凍結特徵方法。GS-MoE 因 per-category MoE "
        "head 超出單卡等級，論文中不列入同級比較。",
    16: "〔Q&A 備用〕reweighting 細節：以視覺流分數離散度比值估計可靠度，"
        "搭配骨架偏移 gate。增益幾乎全部集中在高斯雜訊（平均 +4.83），"
        "其他家族為零＝gate 正確地不出手。已知注意事項：spread proxy 可能過度導流"
        "（SigLIP2-Base 單一種子在高嚴重度高斯最多 −3.2），且方法為 transductive。",
    17: "〔Q&A 備用〕各類別行為：單一種子（seed 42）的探索性分析，非論文正式圖表。",
    18: "〔Q&A 備用〕協定細節：64 幀以下影片排除（UCF 測試集 290 部中評估 254 部，"
        "論文 §4.1 有揭露）；以 validation MIL loss 選 checkpoint；"
        "UCF 影格率較低造成骨架流較弱的可能因素；XD 關閉 smoothness 正則化。",
}

# ---- slides -----------------------------------------------------------------


def s01_title(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_text(s, 1.0, 1.9, SW - 2.0, 1.9,
             [[("Dual-Modal Skeleton-Visual Fusion for Weakly Supervised "
                "Violence Detection:", 30, True, INK)],
              [("A Multi-Backbone Study", 30, True, INK)]],
             align=PP_ALIGN.CENTER, space_after=4)
    add_text(s, 1.0, 4.05, SW - 2.0, 0.5,
             [[("Wei-Han Jeng", 18, True, INK),
               ("   ·   ", 18, False, MUTE),
               ("Chuan-Kai Yang", 18, False, INK)]],
             align=PP_ALIGN.CENTER)
    add_text(s, 1.0, 4.62, SW - 2.0, 0.4,
             [[("National Taiwan University of Science and Technology",
                14, False, MUTE)]], align=PP_ALIGN.CENTER)
    add_text(s, 1.0, 5.55, SW - 2.0, 0.4,
             [[("CGW '26  ·  July 9–10, 2026  ·  Hsinchu, Taiwan"
                "  ·  Paper #32", 14, False, INK)]],
             align=PP_ALIGN.CENTER)
    set_notes(s, NOTES[1])


def s02_motivation(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Violence detection with only video-level labels")
    cards = [
        ("Annotation cost", INK,
         "Frame-level labels for hours of surveillance video are "
         "prohibitively expensive.\n→ Weakly supervised VAD: one label "
         "per video, trained with Multiple Instance Learning (MIL)."),
        ("Appearance isn't enough", BLUE,
         "Strong recent results build on frozen CLIP features — "
         "appearance only.\nFighting, assault, collisions are fundamentally "
         "about motion."),
        ("Deployment drift", RED,
         "Real cameras degrade: noise, blur, compression, lighting.\n"
         "Models are trained and evaluated as if conditions never change."),
    ]
    w, gap, x0, y0 = 3.95, 0.29, 0.45, 1.6
    for i, (head, c, body) in enumerate(cards):
        x = x0 + i * (w + gap)
        add_box(s, x, y0, w, 3.9, WHITE, c, None, line_w=1.5)
        add_text(s, x + 0.25, y0 + 0.3, w - 0.5, 0.6,
                 [[(head, 17, True, c)]])
        paras = [{"runs": [(t, 12.5, False, INK)], "space_after": 8}
                 for t in body.split("\n")]
        add_text(s, x + 0.25, y0 + 1.05, w - 0.5, 2.7, paras,
                 line_spacing=1.08)
    add_text(s, 0.45, 5.9, SW - 0.9, 0.5,
             [[("Can cheap skeleton dynamics make frozen vision-language "
                "features better — and more robust?", 15, True, INK)]],
             align=PP_ALIGN.CENTER)
    add_footer(s, 2)
    set_notes(s, NOTES[2])


def s03_contributions(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Three gaps → three contributions")
    gaps = [
        "Skeleton dynamics + vision-language features: unexplored in "
        "weakly supervised VAD",
        "Sensitivity of fusion to the visual backbone: never "
        "systematically studied",
        "Entropy TTA (TENT/SAR) targets BatchNorm — behavior on "
        "LayerNorm fusion heads unknown",
    ]
    contribs = [
        ("1  Gated dual-modal fusion", ORANGE,
         "Frozen CTR-GCN skeleton + frozen VLM features; only a "
         "lightweight fusion head + MIL head is trained."),
        ("2  Systematic multi-backbone study", BLUE,
         "4 visual backbones × 2 benchmarks × 3 seeds — "
         "CLIP ViT-B/16, SigLIP2 B/16 / SO400M / Giant on UCF-Crime & "
         "XD-Violence."),
        ("3  TTA for frozen dual-modal VAD", PURPLE,
         "Structural null result for entropy TTA + label-free "
         "discriminative-reliability reweighting under corruption."),
    ]
    add_text(s, 0.45, 1.35, 1.5, 0.4, [[("Gaps", 15, True, MUTE)]])
    for i, g in enumerate(gaps):
        add_box(s, 0.45 + i * 4.24, 1.8, 3.95, 1.25, SOFT, MUTE,
                [[(g, 11.5, False, INK)]], line_w=1.0)
    add_text(s, 0.45, 3.4, 3.5, 0.4, [[("Contributions", 15, True, MUTE)]])
    for i, (head, c, body) in enumerate(contribs):
        x = 0.45 + i * 4.24
        add_box(s, x, 3.85, 3.95, 2.5, WHITE, c, None, line_w=1.75)
        add_text(s, x + 0.22, 4.05, 3.55, 0.75, [[(head, 14.5, True, c)]])
        add_text(s, x + 0.22, 4.85, 3.55, 1.4, [[(body, 12, False, INK)]],
                 line_spacing=1.08)
    add_footer(s, 3)
    set_notes(s, NOTES[3])


def s04_architecture(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Architecture — two frozen streams, one trained head")

    # lanes
    add_box(s, 1.9, 1.5, 5.75, 1.7, GREEN_L, None, radius=0.06)
    add_box(s, 1.9, 3.5, 5.75, 1.7, BLUE_L, None, radius=0.06)
    add_text(s, 2.0, 1.55, 4.0, 0.3,
             [[("Skeleton stream — frozen", 10.5, True, GREEN)]])
    add_text(s, 2.0, 3.55, 4.0, 0.3,
             [[("Visual stream — frozen", 10.5, True, BLUE)]])

    # input
    add_box(s, 0.35, 2.95, 1.3, 0.85, WHITE, INK,
            [[("Video frames", 11, True, INK)],
             [("64-frame snippets", 8.5, False, MUTE)]], line_w=1.5)

    # skeleton chain
    add_box(s, 2.15, 2.0, 1.6, 0.95, WHITE, GREEN,
            [[("RTMPose", 11, True, INK)],
             [("COCO-17 kpts, top-2 persons", 8.5, False, MUTE)]])
    add_box(s, 4.05, 2.0, 1.6, 0.95, WHITE, GREEN,
            [[("CTR-GCN", 11, True, INK)],
             [("NTU RGB+D 120, 4 streams", 8.5, False, MUTE)]])
    add_box(s, 5.95, 2.0, 1.55, 0.95, WHITE, GREEN,
            [[("256-d skeleton feature", 9.5, True, GREEN)]],
            dash=MSO_LINE.DASH)

    # visual chain
    add_box(s, 2.15, 4.0, 1.6, 0.95, WHITE, BLUE,
            [[("CLIP / SigLIP2", 11, True, INK)],
             [("frames @ ~1 FPS", 8.5, False, MUTE)]])
    add_box(s, 4.05, 4.0, 1.6, 0.95, WHITE, BLUE,
            [[("mean + max pool", 11, True, INK)],
             [("concat per snippet", 8.5, False, MUTE)]])
    add_box(s, 5.95, 4.0, 1.55, 0.95, WHITE, BLUE,
            [[("1024–3072-d visual feature", 9.5, True, BLUE)]],
            dash=MSO_LINE.DASH)

    # fusion
    add_box(s, 8.0, 2.3, 2.35, 2.35, ORANGE_L, ORANGE,
            [[("Gated Fusion", 13, True, INK)],
             [("(trainable)", 10, True, RED)],
             [("project + LN per modality", 9, False, INK)],
             [("per-dim gate g = σ(·)", 9, False, INK)],
             [("gated blend + residual + LN", 9, False, INK)]], line_w=2.0)

    # prediction
    add_box(s, 10.75, 2.45, 2.2, 0.9, PURPLE_L, PURPLE,
            [[("MIL head", 11.5, True, INK)],
             [("FC + sigmoid, top-k", 8.5, False, MUTE)]], line_w=1.5)
    add_box(s, 10.75, 3.85, 2.2, 0.7, WHITE, PURPLE,
            [[("anomaly score ∈ [0, 1]", 10.5, True, PURPLE)]],
            dash=MSO_LINE.DASH)

    # arrows
    add_arrow(s, 1.65, 3.2, 2.13, 2.6, INK)          # input -> RTMPose
    add_arrow(s, 1.65, 3.6, 2.13, 4.35, INK)         # input -> backbone
    add_arrow(s, 3.75, 2.47, 4.03, 2.47, GREEN)
    add_arrow(s, 5.65, 2.47, 5.93, 2.47, GREEN)
    add_arrow(s, 3.75, 4.47, 4.03, 4.47, BLUE)
    add_arrow(s, 5.65, 4.47, 5.93, 4.47, BLUE)
    add_arrow(s, 7.5, 2.47, 7.98, 3.1, GREEN)        # skel feat -> fusion
    add_arrow(s, 7.5, 4.47, 7.98, 3.9, BLUE)         # vis feat -> fusion
    add_arrow(s, 10.35, 2.9, 10.73, 2.9, ORANGE)     # fusion -> MIL
    add_arrow(s, 11.85, 3.35, 11.85, 3.83, PURPLE)   # MIL -> score

    # training / test-time annotations
    add_box(s, 10.75, 4.9, 2.2, 0.75, None, RED,
            [[("training signal:", 8.5, False, RED)],
             [("MIL ranking loss", 9.5, True, RED)]], dash=MSO_LINE.DASH,
            line_w=1.25)
    add_arrow(s, 11.55, 4.88, 11.55, 4.57, RED, width=1.25,
              dash=MSO_LINE.DASH)
    add_box(s, 7.6, 5.35, 3.1, 0.8, None, RED,
            [[("test-time (slide 11): scale visual stream by w", 9, False,
               RED)],
             [("when its discriminative signal collapses", 9, False, RED)]],
            dash=MSO_LINE.DASH, line_w=1.25)
    add_arrow(s, 9.15, 5.33, 9.15, 4.67, RED, width=1.25, dash=MSO_LINE.DASH)

    add_text(s, 0.45, 6.55, SW - 0.9, 0.45,
             [[("Only the fusion head + MIL head are trained — "
                "0.50–1.02 M parameters · < 5 min per configuration"
                " · single RTX 4090 (24 GB)", 13, True, INK)]],
             align=PP_ALIGN.CENTER)
    add_footer(s, 4)
    set_notes(s, NOTES[4])


def s05_fusion(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Gated fusion & MIL training")
    add_text(s, 0.45, 1.35, 4.0, 0.4,
             [[("Gated fusion (per snippet)", 14, True, ORANGE)]])
    y = 1.85
    for eq in ("eq_proj", "eq_gate", "eq_fuse"):
        h = add_eq(s, eq, 0.7, y, scale=1.22)
        y += h + 0.18
    add_text(s, 0.45, y + 0.05, 4.0, 0.4,
             [[("MIL training objective", 14, True, PURPLE)]])
    add_eq(s, "eq_mil", 0.7, y + 0.5, scale=1.12, max_w=7.15)

    bx = 8.35
    add_box(s, bx - 0.25, 1.5, 5.0, 4.9, SOFT, None, radius=0.05)
    add_text(s, bx, 1.7, 4.6, 4.6, [
        {"runs": [("Reading the gate", 13.5, True, INK)], "space_after": 4},
        {"runs": [("•  residual keeps both modalities — the gate "
                   "modulates their relative contribution", 12, False, INK)],
         "space_after": 10},
        {"runs": [("•  three LayerNorms (no BatchNorm): MIL batches mix "
                   "normal/anomalous bags — this choice matters for TTA "
                   "later", 12, False, INK)], "space_after": 10},
        {"runs": [("Training", 13.5, True, INK)], "space_after": 4,
         "space_before": 6},
        {"runs": [("•  top-k snippet scores per bag (k = 3), margin 1",
                   12, False, INK)], "space_after": 10},
        {"runs": [("•  sparsity + temporal-smoothness regularizers",
                   12, False, INK)], "space_after": 10},
        {"runs": [("Baseline", 13.5, True, INK)], "space_after": 4,
         "space_before": 6},
        {"runs": [("•  late fusion: fixed equal-weight average of two "
                   "per-modality heads", 12, False, INK)], "space_after": 6},
    ], line_spacing=1.05)
    add_footer(s, 5)
    set_notes(s, NOTES[5])


def s06_design(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Study design")
    # dataset cards
    add_box(s, 0.45, 1.5, 5.9, 1.55, WHITE, INK,
            [[("UCF-Crime", 14, True, INK)],
             [("1,900 surveillance videos · 13 anomaly types", 11,
               False, INK)],
             [("metric: frame-level ROC-AUC", 11, True, BLUE)]], line_w=1.5)
    add_box(s, 0.45, 3.25, 5.9, 1.55, WHITE, INK,
            [[("XD-Violence", 14, True, INK)],
             [("4,754 videos (movies, news, online) · 6 violence types",
               11, False, INK)],
             [("metric: frame-level AP", 11, True, BLUE)]], line_w=1.5)
    add_text(s, 0.45, 5.0, 5.9, 0.7,
             [[("† videos < 64 frames excluded — UCF test: 254 of "
                "290 evaluated (disclosed, paper §4.1)", 9.5, False,
                MUTE)]])
    # backbones table
    tb = make_table(s, 6.85, 1.5, 6.0, 1.9, 5, 2, col_widths=[4.0, 2.0])
    rows = [("Visual backbone (frozen)", "feature dim"),
            ("CLIP ViT-B/16 (OpenAI)", "512"),
            ("SigLIP2 ViT-B/16", "768"),
            ("SigLIP2 SO400M", "1152"),
            ("SigLIP2 Giant", "1536")]
    for r, (a, b) in enumerate(rows):
        hdr = r == 0
        set_cell(tb.cell(r, 0), a, size=11, bold=hdr,
                 color=WHITE if hdr else INK, fill=INK if hdr else WHITE,
                 align=PP_ALIGN.LEFT)
        set_cell(tb.cell(r, 1), b, size=11, bold=hdr,
                 color=WHITE if hdr else INK, fill=INK if hdr else WHITE)
    chips = [("all encoders frozen", ORANGE),
             ("seeds {42, 123, 2024} — mean ± std", GREEN),
             ("identical splits & protocol", BLUE),
             ("64-frame snippets · bag T = 32", PURPLE)]
    for i, (t, c) in enumerate(chips):
        add_box(s, 6.85 + (i % 2) * 3.05, 4.15 + (i // 2) * 0.75, 2.9, 0.6,
                WHITE, c, [[(t, 10.5, True, c)]], line_w=1.5)
    add_footer(s, 6)
    set_notes(s, NOTES[6])


def _results_table(s, x, title, header_note, data, hl_row, hl_col):
    add_text(s, x, 1.32, 6.1, 0.35,
             [[(title, 13, True, INK), ("   " + header_note, 10, False,
                                        MUTE)]])
    tb = make_table(s, x, 1.75, 6.15, 2.1, 5, 5,
                    col_widths=[1.75, 1.1, 1.1, 1.1, 1.1])
    heads = ["", "CLIP B/16", "SigLIP2 B/16", "SO400M", "Giant"]
    for c, htxt in enumerate(heads):
        set_cell(tb.cell(0, c), htxt, size=9.5, bold=True, color=WHITE,
                 fill=INK)
    for r, (label, vals) in enumerate(data, start=1):
        band = SOFT if r % 2 == 0 else WHITE
        set_cell(tb.cell(r, 0), label, size=9.5,
                 bold=label.startswith("Gated"), fill=band,
                 align=PP_ALIGN.LEFT)
        if len(vals) == 1:  # skeleton row, backbone-independent
            tb.cell(r, 1).merge(tb.cell(r, 4))
            set_cell(tb.cell(r, 1), vals[0], size=9.5, fill=band, color=INK)
            continue
        for c, v in enumerate(vals, start=1):
            hl = (r == hl_row and c == hl_col)
            set_cell(tb.cell(r, c), v, size=9.5, bold=hl,
                     color=WHITE if hl else INK,
                     fill=ORANGE if hl else band)


def s07_results(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Main results — fusion never hurts, and sets our "
                 "headlines", size=25)
    ucf = [
        ("Skeleton only", ["68.8 ± 2.8   (backbone-independent)"]),
        ("Visual only", ["81.2±0.1", "78.5±0.5", "81.1±0.3",
                         "82.4±0.2"]),
        ("Late fusion", ["78.9±0.4", "77.2±2.0", "79.1±1.3",
                         "79.9±0.8"]),
        ("Gated fusion", ["81.4±0.3", "79.0±0.1", "81.3±0.3",
                          "82.5±0.4"]),
    ]
    xd = [
        ("Skeleton only", ["40.8 ± 0.6   (backbone-independent)"]),
        ("Visual only", ["74.6±1.5", "74.5±2.1", "76.6±2.2",
                         "76.8±1.0"]),
        ("Late fusion", ["63.9±0.5", "63.6±0.7", "65.7±0.5",
                         "65.7±0.4"]),
        ("Gated fusion", ["76.5±0.9", "74.5±0.9", "78.7±0.9",
                          "76.8±2.8"]),
    ]
    _results_table(s, 0.35, "UCF-Crime — AUC (%)",
                   "3-seed mean±std", ucf, hl_row=4, hl_col=4)
    _results_table(s, 6.85, "XD-Violence — AP (%)",
                   "3-seed mean±std", xd, hl_row=4, hl_col=3)
    add_box(s, 2.05, 4.25, 9.2, 0.75, ORANGE_L, ORANGE,
            [[("UCF-Crime  82.5% AUC  (SigLIP2 Giant)      ·      "
               "XD-Violence  78.7% AP  (SigLIP2 SO400M)", 15, True, INK)]],
            line_w=1.75)
    add_text(s, 0.7, 5.3, 12.0, 1.3, [
        {"runs": [("•  Weak alone (68.8 AUC / 40.8 AP), useful together:"
                   " gated fusion ≥ visual-only in 8/8 configurations "
                   "(6/8 strictly; mean +0.6 pp)", 13, False, INK)],
         "space_after": 8},
        {"runs": [("•  Fixed-weight late fusion can hurt — on UCF "
                   "it drags below the visual-only baseline", 13, False,
                   INK)], "space_after": 6},
    ])
    add_footer(s, 7)
    set_notes(s, NOTES[7])


def s08_analysis(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "What the multi-backbone study shows")
    add_picture_fit(s, ASSETS / "fig_backbone_comparison.png", 0, 1.3,
                    w=11.4, center_x=SW / 2)
    y = 1.3 + 11.4 / 3.19 + 0.25  # image height from aspect ratio
    items = [
        ("1", "Complementarity is small but consistent",
         " — no configuration degrades; up to +2.1 AP (XD, SO400M)"),
        ("2", "Gating matters",
         " — beats fixed late fusion by >10 AP on every XD backbone"),
        ("3", "Backbone preference is dataset-specific",
         " — Giant leads UCF; on XD, SO400M wins only under gated "
         "fusion (visual-only backbones within seed variance)"),
    ]
    for i, (n, b, rest) in enumerate(items):
        add_text(s, 0.7, y + i * 0.52, 12.2, 0.5,
                 [[(n + "  ", 13, True, ORANGE), (b, 13, True, INK),
                   (rest, 13, False, INK)]])
    add_footer(s, 8)
    set_notes(s, NOTES[8])


def s09_qualitative(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Qualitative — fusion sees what each stream misses")
    add_picture_fit(s, ASSETS / "fig_temporal_scores_talk.png", 0, 1.4,
                    w=9.6, center_x=SW / 2)
    y = 1.4 + 9.6 / 2.36 + 0.3
    add_text(s, 0.7, y, 12.0, 0.9,
             [[("Fighting test video (UCF-Crime, seed-42 run): gated fusion "
                "stays high through the whole fight — including the opening "
                "the visual stream misses — then drops on the trailing "
                "normal footage; skeleton alone barely reacts. Scores step "
                "at 64-frame-snippet granularity.", 12.5, False, INK)]],
             align=PP_ALIGN.CENTER)
    add_footer(s, 9)
    set_notes(s, NOTES[9])


def s10_sota(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Where this sits — honest comparison")
    add_text(s, 0.45, 1.28, 6.6, 0.4,
             [[("Comparable regime: frozen features · no text branch "
                "· single-GPU head", 11.5, True, MUTE)]])
    rows = [
        ("Method", "Venue", "UCF AUC", "XD AP"),
        ("CLIP-TSA", "ICIP'23", "87.58", "82.19"),
        ("MGFN (I3D)", "AAAI'23", "86.98", "79.19"),
        ("UR-DMU", "AAAI'23", "86.97", "81.66"),
        ("Light-WVAD", "Neurocomp.'24", "84.7", "—"),
        ("RTFM", "ICCV'21", "84.30", "77.81"),
        ("This work", "CGW'26", "82.5", "78.7"),
        ("Sultani et al.", "CVPR'18", "75.41", "—"),
    ]
    tb = make_table(s, 0.45, 1.75, 6.4, 3.6, len(rows), 4,
                    col_widths=[2.3, 1.5, 1.3, 1.3])
    for r, row in enumerate(rows):
        ours = row[0] == "This work"
        for c, v in enumerate(row):
            hdr = r == 0
            set_cell(tb.cell(r, c), v, size=10.5, bold=hdr or ours,
                     color=WHITE if hdr else (RED if ours else INK),
                     fill=INK if hdr else (ORANGE_L if ours else
                                           (SOFT if r % 2 == 0 else WHITE)),
                     align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    add_text(s, 7.25, 1.7, 5.7, 5.0, [
        {"runs": [("No SOTA claim.", 13.5, True, RED),
                  ("  Competitive within this regime — clearest on "
                   "XD AP (on par with RTFM / MGFN); UCF AUC trails the "
                   "modern frozen-feature pack.", 12.5, False, INK)],
         "space_after": 12},
        {"runs": [("Text-aligned / fine-tuned leaders reach ~88–91 UCF "
                   "AUC — the cost of components we deliberately "
                   "excluded.", 12.5, False, INK)], "space_after": 12},
        {"runs": [("Training-free ≠ cheap. ", 13.5, True, INK),
                  ("7B/13B-LLM systems need heavy hardware — EventVAD "
                   "(82.03 / 64.04) an 80 GB A800, LAVAD (80.28 / 62.01) "
                   "dual RTX 3090; one RTX 4090 matches or exceeds them "
                   "here — +14 AP on XD-Violence.", 12.5, False, INK)],
         "space_after": 6},
    ], line_spacing=1.1)
    add_footer(s, 10)
    set_notes(s, NOTES[10])


def s11_tta(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Test-time adaptation — a null result and a fix",
              size=26)
    add_text(s, 0.45, 1.3, 7.4, 5.3, [
        {"runs": [("UCF-Crime-C: ", 13, True, INK),
                  ("4 corruption types × 5 severities × 3 seeds "
                   "(20 conditions)", 13, False, INK)], "space_after": 12},
        {"runs": [("TENT / SAR: ", 13, True, PURPLE),
                  ("adapting the head's 1,536 LayerNorm affine parameters "
                   "moves AUC by ", 13, False, INK),
                  ("< 0.1 pp on every backbone", 13, True, INK),
                  (" — structural: the shift lives in frozen upstream "
                   "features, and AUC is a pure ranking metric.", 13, False,
                   INK)], "space_after": 12},
        {"runs": [("Ours — discriminative-reliability reweighting: ",
                   13, True, ORANGE),
                  ("label-free, tuning-free, adapts zero parameters. When "
                   "the visual stream's score spread collapses, scale it "
                   "down and lean on the corruption-robust skeleton stream.",
                   13, False, INK)], "space_after": 8},
    ], line_spacing=1.08)
    add_eq(s, "eq_w", 0.8, 4.05, scale=1.0)
    add_text(s, 0.45, 4.85, 7.4, 1.7, [
        {"runs": [("+1.21 pp mean", 13.5, True, ORANGE),
                  (" corrupted-AUC · positive on all 4 backbones on "
                   "every seed · up to ", 12.5, False, INK),
                  ("+13.2 pp", 13.5, True, ORANGE),
                  (" where the visual stream collapses hardest (Gaussian, "
                   "severity 5)", 12.5, False, INK)], "space_after": 8},
        {"runs": [("~0 elsewhere — the gate correctly declines to act. "
                   "Rescue, not blanket robustness.", 12.5, True, INK)],
         "space_after": 6},
    ], line_spacing=1.08)
    add_picture_fit(s, ASSETS / "fig_tta_comparison.png", 8.1, 1.75, w=4.9)
    add_footer(s, 11)
    set_notes(s, NOTES[11])


def s12_conclusions(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "Conclusions & limitations")
    findings = [
        ("Skeleton features are weak alone but give a small, consistent "
         "complementary gain on strong visual backbones — and never "
         "degrade them."),
        ("Backbone choice interacts with the dataset and emerges through "
         "fusion — evaluate multi-backbone, not general benchmarks."),
        ("Entropy TTA confined to LN affine updates cannot move our frozen "
         "rank-based detector; reliability reweighting can — when one "
         "modality stays healthy."),
    ]
    for i, f in enumerate(findings):
        add_box(s, 0.45, 1.45 + i * 0.95, 0.55, 0.75, ORANGE, None,
                [[(str(i + 1), 18, True, WHITE)]])
        add_text(s, 1.2, 1.5 + i * 0.95, 11.7, 0.85,
                 [[(f, 13.5, False, INK)]], line_spacing=1.05)
    add_box(s, 0.45, 4.6, 6.1, 1.9, SOFT, None, radius=0.06)
    add_text(s, 0.7, 4.75, 5.7, 1.7, [
        {"runs": [("Limitations", 13, True, RED)], "space_after": 5},
        {"runs": [("frozen design trails fine-tuned / text-aligned leaders "
                   "· reweighting is transductive (whole-condition "
                   "statistics) · 2 datasets, synthetic corruptions",
                   11.5, False, INK)], "space_after": 4},
    ], line_spacing=1.1)
    add_box(s, 6.8, 4.6, 6.1, 1.9, SOFT, None, radius=0.06)
    add_text(s, 7.05, 4.75, 5.7, 1.7, [
        {"runs": [("Modular & efficient", 13, True, GREEN)], "space_after": 5},
        {"runs": [("swap in a new VLM by re-extracting features — the "
                   "0.5–1.0 M-parameter head retrains in < 5 minutes on "
                   "one RTX 4090", 11.5, False, INK)], "space_after": 4},
    ], line_spacing=1.1)
    add_footer(s, 12)
    set_notes(s, NOTES[12])


def s13_thanks(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_text(s, 1.0, 2.0, SW - 2.0, 1.0,
             [[("Thank you — Questions?", 36, True, INK)]],
             align=PP_ALIGN.CENTER)
    chips = [("UCF-Crime 82.5% AUC", BLUE),
             ("XD-Violence 78.7% AP", GREEN),
             ("corrupted-AUC +1.21 pp (mean)", ORANGE)]
    w = 3.6
    total = 3 * w + 2 * 0.4
    x0 = (SW - total) / 2
    for i, (t, c) in enumerate(chips):
        add_box(s, x0 + i * (w + 0.4), 3.6, w, 0.7, WHITE, c,
                [[(t, 13.5, True, c)]], line_w=1.75)
    add_text(s, 1.0, 4.9, SW - 2.0, 0.4,
             [[("Wei-Han Jeng · austin60616@gmail.com · code & "
                "experiment configurations to be released", 13, False,
                MUTE)]], align=PP_ALIGN.CENTER)
    add_footer(s, 13)
    set_notes(s, NOTES[13])


def s14_divider(prs):
    s = blank_slide(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW),
                            Inches(SH))
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    bg.shadow.inherit = False
    add_text(s, 1.0, 3.1, SW - 2.0, 1.0,
             [[("Backup — Q&A", 34, True, WHITE)]],
             align=PP_ALIGN.CENTER)
    set_notes(s, NOTES[14])


def s15_full_sota(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "B1 · Full comparison incl. heavier-budget regimes",
              size=24)
    rows = [
        ("Method", "Venue", "UCF AUC", "XD AP", "Setting"),
        ("PI-VAD", "CVPR'25", "90.33", "85.37",
         "frozen I3D; +text, +5 train-time modalities"),
        ("DSANet", "AAAI'26", "89.44", "86.95",
         "frozen CLIP; +text alignment"),
        ("VadCLIP", "AAAI'24", "88.02", "84.51",
         "frozen CLIP; +text alignment"),
        ("EventVAD", "ACM MM'25", "82.03", "64.04",
         "training-free 7B MLLM; 80 GB A800"),
        ("LAVAD", "CVPR'24", "80.28", "62.01",
         "training-free 13B LLM; dual RTX 3090"),
        ("CLIP-TSA", "ICIP'23", "87.58", "82.19", "frozen CLIP; no text"),
        ("MGFN (I3D)", "AAAI'23", "86.98", "79.19", "frozen I3D; no text"),
        ("UR-DMU", "AAAI'23", "86.97", "81.66", "frozen I3D; no text"),
        ("Light-WVAD", "Neurocomp.'24", "84.7", "—",
         "frozen I3D; no text"),
        ("RTFM", "ICCV'21", "84.30", "77.81", "frozen I3D; no text"),
        ("This work", "CGW'26", "82.5", "78.7",
         "frozen skeleton+CLIP/SigLIP2; no text"),
        ("Sultani et al.", "CVPR'18", "75.41", "—",
         "frozen C3D; no text"),
    ]
    tb = make_table(s, 0.45, 1.35, 12.4, 4.7, len(rows), 5,
                    col_widths=[2.0, 1.5, 1.2, 1.2, 6.5])
    for r, row in enumerate(rows):
        ours = row[0] == "This work"
        for c, v in enumerate(row):
            hdr = r == 0
            set_cell(tb.cell(r, c), v, size=9, bold=hdr or ours,
                     color=WHITE if hdr else (RED if ours else INK),
                     fill=INK if hdr else (ORANGE_L if ours else
                                           (SOFT if r % 2 == 0 else WHITE)),
                     align=PP_ALIGN.LEFT if c in (0, 4) else
                     PP_ALIGN.CENTER)
    add_text(s, 0.45, 6.25, 12.4, 0.6,
             [[("Rows 1–5: heavier budget (text alignment, extra "
                "modalities, MLLM). Rows 6–12: comparable regime. "
                "GS-MoE (ICCV'25, 91.58/82.89) omitted from the comparable "
                "block: per-category MoE head exceeds the single-GPU-class "
                "regime.", 9.5, False, MUTE)]])
    add_footer(s, 15)
    set_notes(s, NOTES[15])


def s16_tta_breakdown(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "B2 · TTA breakdown & method detail", size=24)
    add_text(s, 0.45, 1.3, 7.0, 0.4,
             [[("Reweighting: visual-score spread ratio + skeleton-shift "
                "gate", 12, True, INK)]])
    add_eq(s, "eq_w", 0.7, 1.75, scale=1.05)
    rows = [
        ("Backbone", "Gaussian", "Motion blur", "JPEG", "Brightness"),
        ("CLIP ViT-B/16", "+2.67", "0.00", "0.00", "+0.01"),
        ("SigLIP2 B/16", "+5.99", "0.00", "0.00", "0.00"),
        ("SigLIP2 SO400M", "+8.95", "0.00", "0.00", "0.00"),
        ("SigLIP2 Giant", "+1.69", "0.00", "0.00", "0.00"),
        ("Average", "+4.83", "0.00", "0.00", "0.00"),
    ]
    tb = make_table(s, 0.45, 2.6, 7.4, 2.4, len(rows), 5,
                    col_widths=[2.2, 1.3, 1.3, 1.3, 1.3])
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            hdr = r == 0
            avg = row[0] == "Average"
            set_cell(tb.cell(r, c), v, size=10, bold=hdr or avg,
                     color=WHITE if hdr else INK,
                     fill=INK if hdr else (ORANGE_L if avg else
                                           (SOFT if r % 2 == 0 else WHITE)),
                     align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    add_text(s, 0.45, 5.25, 7.4, 0.4,
             [[("Gain by corruption family — 3-seed mean ΔAUC, "
                "averaged over 5 severities (paper breakdown table)", 9.5,
                False, MUTE)]])
    add_text(s, 8.2, 1.75, 4.8, 4.6, [
        {"runs": [("Zeros are a feature", 13, True, GREEN)], "space_after": 4},
        {"runs": [("visual still reliable, or skeleton also corrupted "
                   "→ the gate declines to act", 11.5, False, INK)],
         "space_after": 12},
        {"runs": [("Known caveat", 13, True, RED)], "space_after": 4},
        {"runs": [("the spread proxy can over-route: one seed of SigLIP2-"
                   "B/16 loses up to 3.2 pp at high-severity Gaussian "
                   "(3-seed average still positive)", 11.5, False, INK)],
         "space_after": 12},
        {"runs": [("Transductive", 13, True, PURPLE)], "space_after": 4},
        {"runs": [("statistics pooled over a whole corruption condition; "
                   "streaming variant = future work", 11.5, False, INK)],
         "space_after": 6},
    ], line_spacing=1.08)
    add_footer(s, 16)
    set_notes(s, NOTES[16])


def s17_per_category(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "B3 · Per-category behavior (exploratory)", size=24)
    add_picture_fit(s, ASSETS / "fig_gating_distribution.png", 0, 1.5,
                    w=11.6, center_x=SW / 2)
    y = 1.5 + 11.6 / 2.90 + 0.25
    add_text(s, 0.7, y, 12.0, 0.8,
             [[("Exploratory single-run breakdown (seed 42) — not a "
                "paper figure; paper tables are 3-seed. XD categories: "
                "Fighting, Shooting, Riot, Abuse, Car Accident, Explosion.",
                11, False, MUTE)]], align=PP_ALIGN.CENTER)
    add_footer(s, 17)
    set_notes(s, NOTES[17])


def s18_protocol(prs):
    s = blank_slide(prs)
    add_strip(s)
    add_title(s, "B4 · Protocol details", size=24)
    items = [
        ("Exclusions", "videos < 64 frames excluded (172 UCF, 1 XD); UCF "
         "test: 254 of 290 evaluated — disclosed in paper §4.1"),
        ("Model selection", "15% stratified validation split; checkpoint "
         "chosen by validation MIL loss (no frame-level peeking)"),
        ("Frame-rate asymmetry", "UCF distributed at ~3 FPS (64-frame "
         "snippet ≈ 20 s); XD native ~24 FPS (≈ 2.7 s) — a "
         "plausible factor in the weaker UCF skeleton-only result"),
        ("Snippet → frame expansion", "UCF: final-snippet score "
         "repeated to the full frame grid; XD: evaluated on the "
         "snippet-aligned portion"),
        ("Smoothness λ₂", "8×10⁻⁴ on UCF-Crime; 0 "
         "on XD-Violence (over-regularizes heterogeneous, rapidly-cut "
         "sources)"),
        ("Optimization", "AdamW, lr 1×10⁻⁴, 5-epoch warmup + "
         "cosine; batch 16 normal/anomalous pairs; bag T = 32; top-k with "
         "k = 3"),
    ]
    for i, (head, body) in enumerate(items):
        y = 1.4 + i * 0.87
        add_text(s, 0.45, y, 3.0, 0.8, [[(head, 12.5, True, BLUE)]])
        add_text(s, 3.6, y, 9.3, 0.8, [[(body, 11.5, False, INK)]],
                 line_spacing=1.02)
    add_footer(s, 18)
    set_notes(s, NOTES[18])


# ---- main -------------------------------------------------------------------

def build():
    print("[1/3] Converting paper figures...")
    convert_figures()
    print("[2/3] Rendering equations + talk figure...")
    render_equations()
    render_temporal_talk_figure()
    print("[3/3] Building deck...")
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = (
        "Dual-Modal Skeleton-Visual Fusion for Weakly Supervised Violence "
        "Detection: A Multi-Backbone Study")
    prs.core_properties.author = "Wei-Han Jeng"

    for fn in (s01_title, s02_motivation, s03_contributions,
               s04_architecture, s05_fusion, s06_design, s07_results,
               s08_analysis, s09_qualitative, s10_sota, s11_tta,
               s12_conclusions, s13_thanks, s14_divider, s15_full_sota,
               s16_tta_breakdown, s17_per_category, s18_protocol):
        fn(prs)

    prs.save(OUT)
    print(f"Saved: {OUT}")

    # sanity: re-open, count slides and notes coverage
    check = Presentation(OUT)
    n_notes = sum(
        1 for sl in check.slides
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip()
    )
    n_slides = len(check.slides._sldIdLst)
    print(f"Re-opened OK: {n_slides} slides, {n_notes} with speaker notes")
    assert n_slides == 18, f"expected 18 slides, got {n_slides}"


if __name__ == "__main__":
    build()
